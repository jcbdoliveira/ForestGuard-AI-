"""
Catálogo MapBiomas Coleção 11 — área de floresta por município (ha).

Fonte:
  data/mapbiomas/MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.xlsx
  (espelho: https://github.com/jcbdoliveira/MAPBiomas/raw/refs/heads/main/MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.xlsx)

Índice pré-computado (rápido):
  data/mapbiomas/forest_by_municipality.json

A série de perda anual é a diferença ano-a-ano da soma das classes
"1. Forest" (formação florestal, savana, floresta alagável, etc.).
Valores em hectares. Não substitui PRODES oficial para fins legais.
"""

from __future__ import annotations

import json
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Optional

ROOT = Path(__file__).resolve().parents[3]
INDEX_PATH = ROOT / "data" / "mapbiomas" / "forest_by_municipality.json"
XLSX_PATH = ROOT / "data" / "mapbiomas" / "MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.xlsx"
SOURCE_URL = (
    "https://github.com/jcbdoliveira/MAPBiomas/raw/refs/heads/main/"
    "MAPBIOMAS_BRAZIL-COL.11-BIOME_STATE_MUNICIPALITY.xlsx"
)


@lru_cache(maxsize=1)
def load_forest_index() -> Dict[str, Any]:
    if not INDEX_PATH.exists():
        return {}
    return json.loads(INDEX_PATH.read_text(encoding="utf-8"))


def get_mapbiomas_municipality(ibge_code: str) -> Optional[Dict[str, Any]]:
    """Retorna estatísticas MapBiomas para o código IBGE (7 dígitos) ou None."""
    code = str(ibge_code).strip()
    idx = load_forest_index()
    rec = idx.get(code)
    if not rec:
        return None
    return dict(rec)


def mapbiomas_series_for_prodes(ibge_code: str) -> Optional[Dict[str, Any]]:
    """
    Converte registro MapBiomas no formato usado pelo catálogo PRODES:
      historic_prodes_deforestation_ha, forest_remnant_ha, mean_annual_deforestation_rate_pct
    """
    rec = get_mapbiomas_municipality(ibge_code)
    if not rec:
        return None
    loss = rec.get("approx_annual_forest_loss_ha") or {}
    remnant = float(rec.get("forest_remnant_ha") or 0)
    if not loss:
        return None
    vals = [float(v) for v in loss.values() if float(v) >= 0]
    mean_loss = sum(vals) / len(vals) if vals else 0.0
    rate_pct = (mean_loss / remnant * 100.0) if remnant > 0 else 0.0
    return {
        "historic_prodes_deforestation_ha": {str(k): float(v) for k, v in loss.items()},
        "forest_remnant_ha": remnant,
        "forest_ha_by_year": rec.get("forest_ha_by_year") or {},
        "mean_annual_deforestation_rate_pct": round(rate_pct, 4),
        "biome": rec.get("biome"),
        "name": rec.get("name"),
        "state": rec.get("state"),
        "source": rec.get("source") or "MapBiomas Coleção 11",
        "source_url": rec.get("source_url") or SOURCE_URL,
    }


def rebuild_index_from_xlsx(xlsx_path: Optional[Path] = None) -> int:
    """Reconstrói forest_by_municipality.json a partir do XLSX. Retorna nº de municípios."""
    from collections import defaultdict
    import openpyxl

    path = Path(xlsx_path) if xlsx_path else XLSX_PATH
    if not path.exists():
        raise FileNotFoundError(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    year_cols = [
        (i, h)
        for i, h in enumerate(header)
        if isinstance(h, str) and h.startswith("y") and h[1:].isdigit()
    ]
    forest: Dict[str, Dict[str, float]] = defaultdict(lambda: defaultdict(float))
    meta: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        code = str(r[5] or "").strip()
        if not code:
            continue
        cl1 = str(r[10] or "")
        if not cl1.startswith("1. Forest"):
            continue
        meta[code] = {"name": r[6], "state": r[4], "biome": r[2]}
        for i, h in year_cols:
            try:
                v = float(r[i] or 0)
            except Exception:
                v = 0.0
            forest[code][h[1:]] += v
    out: Dict[str, Any] = {}
    for code, years in forest.items():
        series = {y: round(years[y], 2) for y in sorted(years) if 2015 <= int(y) <= 2024}
        if not series:
            continue
        last = sorted(series.keys())[-1]
        ys = sorted(series.keys())
        loss: Dict[str, float] = {}
        for a, b in zip(ys, ys[1:]):
            loss[b] = round(max(0.0, series[a] - series[b]), 2)
        out[code] = {
            **meta.get(code, {}),
            "forest_ha_by_year": series,
            "forest_remnant_ha": series[last],
            "approx_annual_forest_loss_ha": loss,
            "source": "MapBiomas Coleção 11 (soma classes 1. Forest)",
            "source_url": SOURCE_URL,
        }
    INDEX_PATH.parent.mkdir(parents=True, exist_ok=True)
    INDEX_PATH.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    load_forest_index.cache_clear()
    return len(out)
